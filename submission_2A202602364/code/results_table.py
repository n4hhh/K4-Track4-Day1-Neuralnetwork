"""Luu JSON va tao bang experiments.xlsx."""

from __future__ import annotations

import json
import math
from pathlib import Path

import numpy as np
import openpyxl
from openpyxl.formula.translate import Translator

FORMULA_COLUMNS = {
    "step0_gap_vs_lnC",
    "gap_val_minus_train",
    "delta_val_f1_vs_base",
    "beyond_noise",
}


def _json_safe(value):
    """Doi NumPy va gia tri dac biet sang JSON."""
    if isinstance(value, dict):
        return {key: _json_safe(item) for key, item in value.items()}

    if isinstance(value, (list, tuple)):
        return [_json_safe(item) for item in value]

    if isinstance(value, np.generic):
        value = value.item()

    if isinstance(value, float) and not math.isfinite(value):
        return None

    if isinstance(value, Path):
        return str(value)

    return value


def save_result(
    result: dict,
    results_dir: str = "../results",
) -> str:
    """Luu cfg, history va summary; khong luu best_state."""
    exp_id = result["cfg"]["exp_id"]

    if Path(exp_id).name != exp_id:
        raise ValueError("exp_id khong hop le")

    output_dir = Path(results_dir)
    output_dir.mkdir(
        parents=True,
        exist_ok=True,
    )

    output_path = output_dir / f"{exp_id}.json"

    payload = {
        "cfg": result["cfg"],
        "history": result["history"],
        "summary": result["summary"],
    }

    with output_path.open(
        "w",
        encoding="utf-8",
    ) as file:
        json.dump(
            _json_safe(payload),
            file,
            ensure_ascii=False,
            indent=2,
        )

    print("Saved result:", output_path)
    return str(output_path)


def load_results(
    results_dir: str = "../results",
) -> list[dict]:
    """Doc JSON thi nghiem, bo qua cac JSON thong ke hoac metadata."""
    input_dir = Path(results_dir)

    if not input_dir.exists():
        return []

    results = []

    for path in sorted(input_dir.glob("*.json")):
        with path.open("r", encoding="utf-8") as file:
            payload = json.load(file)

        if not isinstance(payload, dict):
            continue

        cfg = payload.get("cfg")
        history = payload.get("history")
        summary = payload.get("summary")

        if not all(isinstance(item, dict) for item in (cfg, history, summary)):
            continue

        exp_id = cfg.get("exp_id")

        if not isinstance(exp_id, str) or not exp_id:
            continue

        results.append(payload)

    results.sort(key=lambda item: item["cfg"]["exp_id"])
    return results


def to_row(
    result: dict,
    eval_scores: dict | None = None,
    notes: str = "",
) -> dict:
    """Chuyen ket qua thi nghiem thanh mot dong Excel."""
    cfg = result["cfg"]
    summary = result["summary"]
    exp_id = cfg["exp_id"]

    hidden = "x".join(str(width) for width in cfg["hidden"])

    eval_acc = None
    eval_macro_f1 = None

    if eval_scores is not None:
        eval_acc = eval_scores.get(
            "accuracy",
            eval_scores.get("eval_acc"),
        )
        eval_macro_f1 = eval_scores.get(
            "macro_f1",
            eval_scores.get("eval_macro_f1"),
        )

    return {
        "exp_id": exp_id,
        "group": cfg.get("group"),
        "description": cfg.get("description"),
        "loss": cfg.get("loss"),
        "optimizer": cfg.get("optimizer"),
        "lr": cfg.get("lr"),
        "weight_decay": cfg.get("weight_decay"),
        "batch": cfg.get("batch"),
        "epochs": cfg.get("epochs"),
        "hidden": hidden,
        "dropout": cfg.get("dropout"),
        "clip_norm": cfg.get("clip_norm"),
        "precision": cfg.get("precision"),
        "init": cfg.get("init"),
        "seed": cfg.get("seed"),
        "step0_loss": summary.get("step0_loss"),
        "best_val_loss": summary.get("best_val_loss"),
        "best_epoch": summary.get("best_epoch"),
        "final_train_loss": summary.get("final_train_loss"),
        "final_val_loss": summary.get("final_val_loss"),
        "val_acc": summary.get("val_acc"),
        "val_macro_f1": summary.get("val_macro_f1"),
        "time_per_epoch_s": summary.get("time_per_epoch_s"),
        "peak_mem_MB": summary.get("peak_mem_MB"),
        "diverged": summary.get("diverged"),
        "eval_acc": eval_acc,
        "eval_macro_f1": eval_macro_f1,
        "figure_file": f"figures/{exp_id}.png",
        "notes": notes,
    }


def write_xlsx(
    rows: list[dict],
    template_path: str,
    out_path: str,
) -> None:
    """Dien cac dong vao sheet Experiments."""
    workbook = openpyxl.load_workbook(
        template_path,
        data_only=False,
    )

    if "Experiments" not in workbook.sheetnames:
        raise KeyError("Template khong co sheet Experiments")

    sheet = workbook["Experiments"]

    headers = {cell.value: cell.column for cell in sheet[1] if cell.value is not None}

    if "exp_id" not in headers:
        raise KeyError("Template khong co cot exp_id")

    formula_templates = {}

    for name in FORMULA_COLUMNS:
        column = headers.get(name)

        if column is None:
            continue

        value = sheet.cell(
            row=2,
            column=column,
        ).value

        if isinstance(value, str) and value.startswith("="):
            formula_templates[column] = value

    for row_index, row in enumerate(
        rows,
        start=2,
    ):
        for key, value in row.items():
            if key in FORMULA_COLUMNS:
                continue

            column = headers.get(key)

            if column is not None:
                sheet.cell(
                    row=row_index,
                    column=column,
                    value=_json_safe(value),
                )

        if row_index > 2:
            for column, formula in formula_templates.items():
                source = sheet.cell(
                    row=2,
                    column=column,
                )
                destination = sheet.cell(
                    row=row_index,
                    column=column,
                )

                destination.value = Translator(
                    formula,
                    origin=source.coordinate,
                ).translate_formula(destination.coordinate)

    output_path = Path(out_path)
    output_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    workbook.save(output_path)

    print("Saved experiment table:", output_path)
